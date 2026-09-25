"""读取 Excel / CSV，统一成 list[list] 的行表结构。"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional, Sequence

EXCEL_EXT = {'.xlsx', '.xlsm'}
CSV_EXT = {'.csv'}
SUPPORTED_EXT = EXCEL_EXT | CSV_EXT | {'.xls'}

# 单元格自带类型信息的扩展名：xlsx/xls 的单元格有类型（数字就是数字，文本就是文本）
TYPED_EXT = EXCEL_EXT | {'.xls'}

# 编码尝试顺序：中文环境最常见的是 gb18030（gbk 超集）与 utf-8-sig（Excel 导出）
CSV_ENCODINGS = ('utf-8-sig', 'gb18030', 'utf-8', 'gbk')
CSV_DELIMITERS = (',', ';', '\t', '|')


def has_type_info(path) -> bool:
    """该文件的数据是否自带单元格类型信息。

    xlsx/xls 里数字与文本是分开存的，读出来就带类型；
    CSV 是无类型纯文本，**每一格读出来都是 str** —— 要不要按数字处理
    只能靠内容推断（这就是 `infer_types = auto` 只对 CSV 生效的原因：
    Excel 源里写成文本的单元格，是用户有意写成文本的，不该被改写）。
    """
    return Path(path).suffix.lower() in TYPED_EXT


@dataclass
class Sheet:
    name: str
    rows: List[List[object]] = field(default_factory=list)
    ncols: int = 0

    @property
    def nrows(self) -> int:
        return len(self.rows)


class ReadError(Exception):
    """读取失败（文件格式、依赖缺失、编码等）。"""


def _clean(rows: Sequence[Sequence[object]], na_to_none: bool = True) -> List[List[object]]:
    """统一空值语义：Excel 的 None、pandas 的 NaN 都归一到 None。"""
    out: List[List[object]] = []
    for r in rows:
        row = list(r)
        if na_to_none:
            for i, v in enumerate(row):
                if v is None:
                    continue
                if isinstance(v, float) and v != v:        # NaN
                    row[i] = None
        out.append(row)
    return out


def _read_csv(path: Path, encoding: Optional[str], delimiter: Optional[str]) -> List[Sheet]:
    encodings = (encoding,) if encoding else CSV_ENCODINGS
    last_err: Optional[Exception] = None
    for enc in encodings:
        try:
            text = path.read_text(encoding=enc)
        except (UnicodeDecodeError, LookupError) as e:
            last_err = e
            continue
        seps = (delimiter,) if delimiter else CSV_DELIMITERS
        best = None                                    # (列数, 行数据) —— 取切得最细的分隔符
        for sep in seps:
            try:
                rows = list(csv.reader(io.StringIO(text), delimiter=sep))
            except csv.Error as e:
                last_err = e
                continue
            ncols = max((len(r) for r in rows), default=0)
            if best is None or ncols > best[0]:
                best = (ncols, rows)
        if best is not None:
            return [Sheet(path.stem, _clean(best[1]), best[0])]
    raise ReadError('CSV 解析失败（已尝试编码：{}）：{}'.format('/'.join(encodings), last_err))


def _read_excel(path: Path) -> List[Sheet]:
    try:
        import openpyxl
    except ImportError:
        raise ReadError('缺少依赖 openpyxl，请执行：pip install openpyxl')

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        sheets = []
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            sheets.append(Sheet(ws.title, _clean(rows), ws.max_column or 0))
        return sheets
    finally:
        wb.close()


def _read_xls(path: Path) -> List[Sheet]:
    try:
        import pandas as pd
    except ImportError:
        raise ReadError('.xls 需要额外依赖，请执行：pip install "excel2sql[xls]"（或先另存为 .xlsx）')

    frames = pd.read_excel(path, sheet_name=None, header=None)
    sheets = []
    for name, df in frames.items():
        rows = [[None if pd.isna(c) else c for c in row] for row in df.values.tolist()]
        sheets.append(Sheet(str(name), rows, len(df.columns)))
    return sheets


def read_sheets(path, encoding: Optional[str] = None, delimiter: Optional[str] = None) -> List[Sheet]:
    """读取文件，返回全部工作表。"""
    path = Path(path)
    if not path.is_file():
        raise ReadError('文件不存在：{}'.format(path))
    ext = path.suffix.lower()
    if ext in CSV_EXT:
        return _read_csv(path, encoding, delimiter)
    if ext in EXCEL_EXT:
        return _read_excel(path)
    if ext == '.xls':
        return _read_xls(path)
    raise ReadError('不支持的文件类型：{}（支持 .xlsx/.xlsm/.xls/.csv）'.format(ext or path.name))


@dataclass
class SheetInfo:
    """轻量探测结果：只读工作簿元信息，不遍历单元格。"""
    name: str
    nrows: int = -1                 # -1 表示未知（odf/dimension 缺失时）
    ncols: int = -1

    def shape_text(self) -> str:
        rows = '?' if self.nrows < 0 else '{:,}'.format(self.nrows)
        cols = '?' if self.ncols < 0 else str(self.ncols)
        return '{} 行 x {} 列'.format(rows, cols)


def _open_excel(path):
    try:
        import openpyxl
    except ImportError:
        raise ReadError('缺少依赖 openpyxl，请执行：pip install openpyxl')
    return openpyxl.load_workbook(path, data_only=True, read_only=True)


def probe_sheets(path) -> List[SheetInfo]:
    """列出 sheet 名与规模，成本远低于把数据读进内存。

    大表（几万行）下这一步通常 < 1s，而全量读取可能几十秒，
    因此交互流程应先 probe、再按选中的 sheet 读取。
    """
    path = Path(path)
    if not path.is_file():
        raise ReadError('文件不存在：{}'.format(path))
    ext = path.suffix.lower()

    if ext in EXCEL_EXT:
        wb = _open_excel(path)
        try:
            return [SheetInfo(ws.title, ws.max_row or -1, ws.max_column or -1) for ws in wb.worksheets]
        finally:
            wb.close()
    # CSV / xls 读取本身很快，直接读
    return [SheetInfo(s.name, s.nrows, max((len(r) for r in s.rows), default=0))
            for s in read_sheets(path)]


def read_sheet_rows(path, sheet: Optional[str] = None,
                    encoding: Optional[str] = None, delimiter: Optional[str] = None) -> Sheet:
    """只读取指定 sheet 的数据（xlsx 下不会遍历其他 sheet）。"""
    path = Path(path)
    ext = path.suffix.lower()

    if ext in EXCEL_EXT:
        wb = _open_excel(path)
        try:
            names = [ws.title for ws in wb.worksheets]
            if sheet and sheet not in names:
                raise ReadError('找不到 sheet：{}（可选：{}）'.format(sheet, ', '.join(names)))
            ws = wb[sheet] if sheet else wb.worksheets[0]
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            return Sheet(ws.title, _clean(rows), ws.max_column or 0)
        finally:
            wb.close()

    sheets = read_sheets(path, encoding=encoding, delimiter=delimiter)
    if sheet:
        for s in sheets:
            if s.name == sheet:
                return s
        raise ReadError('找不到 sheet：{}（可选：{}）'.format(sheet, ', '.join(s.name for s in sheets)))
    return sheets[0]


def scan_dir(directory) -> List[Path]:
    """列出目录下受支持的数据文件（忽略 Excel 临时文件 ~$xxx）。"""
    d = Path(directory)
    if not d.is_dir():
        return []
    return sorted(p for p in d.iterdir()
                  if p.is_file() and p.suffix.lower() in SUPPORTED_EXT and not p.name.startswith('~$'))
